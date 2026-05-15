"""Tests for reports.template — placeholder detection and fill logic."""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from localdocsai.reports.template import (
    _classify,
    _extract_label,
    _find_preceding_heading,
    detect_fields,
    fill_and_save,
)

# ---------------------------------------------------------------------------
# _classify
# ---------------------------------------------------------------------------


def test_classify_date() -> None:
    assert _classify("fecha de emisión") == "date"


def test_classify_title() -> None:
    assert _classify("Proyecto / Referencia") == "title"


def test_classify_topic() -> None:
    assert _classify("Tema / Área") == "topic"


def test_classify_objective() -> None:
    assert _classify("Objetivo de la Consulta") == "objective"


def test_classify_synthesis() -> None:
    assert _classify("Síntesis de la Información") == "synthesis"


def test_classify_unknown() -> None:
    assert _classify("Número de expediente") == "unknown"


# ---------------------------------------------------------------------------
# _extract_label
# ---------------------------------------------------------------------------


def test_extract_label_bold_prefix() -> None:
    assert _extract_label("*Fecha:* ", "fallback") == "Fecha"


def test_extract_label_heading() -> None:
    assert _extract_label("## 1. Objetivo de la Consulta", "fallback") == "Objetivo de la Consulta"


def test_extract_label_plain() -> None:
    label = _extract_label("Algún texto sin markdown", "fallback")
    assert "Algún texto" in label


def test_extract_label_fallback_on_empty() -> None:
    label = _extract_label("", "Campo 1")
    assert label == "Campo 1"


# ---------------------------------------------------------------------------
# _find_preceding_heading
# ---------------------------------------------------------------------------


def test_find_preceding_heading_skips_empty() -> None:
    texts = ["## Sección", "", ""]
    result = _find_preceding_heading(texts, 2)
    assert result == "## Sección"


def test_find_preceding_heading_skips_placeholder_para() -> None:
    texts = ["## Objetivo", "[placeholder aquí]"]
    result = _find_preceding_heading(texts, 1)
    assert result == "## Objetivo"


def test_find_preceding_heading_no_heading() -> None:
    result = _find_preceding_heading([], 0)
    assert result == ""


# ---------------------------------------------------------------------------
# detect_fields — integration test against the default template
# ---------------------------------------------------------------------------

_TEMPLATE = (
    Path(__file__).parent.parent.parent
    / "resources"
    / "report_templates"
    / "nota_tecnica.docx"
)


@pytest.mark.skipif(not _TEMPLATE.exists(), reason="template not present in CI")
def test_detect_fields_finds_five_fields() -> None:
    fields = detect_fields(_TEMPLATE)
    assert len(fields) == 5


@pytest.mark.skipif(not _TEMPLATE.exists(), reason="template not present in CI")
def test_detect_fields_semantic_types() -> None:
    fields = detect_fields(_TEMPLATE)
    types = {f.semantic_type for f in fields}
    assert "date" in types
    assert "title" in types
    assert "topic" in types
    assert "objective" in types
    assert "synthesis" in types


@pytest.mark.skipif(not _TEMPLATE.exists(), reason="template not present in CI")
def test_detect_fields_labels_are_readable() -> None:
    fields = detect_fields(_TEMPLATE)
    for f in fields:
        assert len(f.label) > 0
        assert not f.label.startswith("[")  # should not be the raw placeholder


# ---------------------------------------------------------------------------
# fill_and_save — integration test
# ---------------------------------------------------------------------------


@pytest.mark.skipif(not _TEMPLATE.exists(), reason="template not present in CI")
def test_fill_and_save_replaces_placeholders(tmp_path: Path) -> None:
    fields = detect_fields(_TEMPLATE)
    for f in fields:
        object.__setattr__(f, "value", f"VALOR_{f.label.upper()[:10]}")

    out = tmp_path / "output.docx"
    fill_and_save(_TEMPLATE, fields, out)

    assert out.exists()
    from docx import Document

    doc = Document(str(out))
    full = " ".join("".join(r.text for r in p.runs) for p in doc.paragraphs)
    # All original bracket patterns should be gone
    assert not re.search(r"\[DD/MM", full)
    assert not re.search(r"\[Ej\.", full)
